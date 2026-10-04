import io
import pytest
import fitz
import docx
from app.services.extractor import (
    extract_document,
    extract_pdf,
    extract_docx,
    extract_txt,
    DocumentExtractionError,
    _count_words_and_sentences
)

def test_count_words_and_sentences():
    text = "Natural Language Processing is fascinating. Look inside. Understand more!"
    chars, words, sents = _count_words_and_sentences(text)
    assert chars == len(text)
    assert words == 9
    assert sents == 3

def test_extract_txt_utf8():
    sample_text = "LookinDoc is an NLP-powered document intelligence platform. It analyzes documents accurately."
    file_bytes = sample_text.encode("utf-8")
    doc = extract_document(file_bytes, "sample.txt")
    assert doc.file_type == "txt"
    assert "LookinDoc is an NLP-powered" in doc.extracted_text
    assert doc.word_count > 0
    assert doc.sentence_count == 2
    assert doc.page_count >= 1

def test_extract_txt_latin1():
    sample_text = "Document with special characters: café, résumé, naïve."
    file_bytes = sample_text.encode("latin-1")
    doc = extract_document(file_bytes, "resume.txt")
    assert doc.file_type == "txt"
    assert "café" in doc.extracted_text

def test_extract_pdf():
    # Create in-memory PDF using PyMuPDF
    pdf_doc = fitz.open()
    page = pdf_doc.new_page()
    page.insert_text((50, 72), "LookinDoc PDF Extraction Test. PyMuPDF extracts text perfectly.")
    pdf_bytes = pdf_doc.write()
    pdf_doc.close()

    doc = extract_document(pdf_bytes, "test_doc.pdf")
    assert doc.file_type == "pdf"
    assert doc.page_count == 1
    assert "LookinDoc PDF Extraction Test" in doc.extracted_text
    assert doc.word_count > 5

def test_extract_docx():
    # Create in-memory DOCX using python-docx
    docx_doc = docx.Document()
    docx_doc.add_heading("LookinDoc DOCX Test", level=1)
    docx_doc.add_paragraph("This paragraph contains natural language text for testing DOCX extraction.")
    
    stream = io.BytesIO()
    docx_doc.save(stream)
    docx_bytes = stream.getvalue()

    doc = extract_document(docx_bytes, "test_doc.docx")
    assert doc.file_type == "docx"
    assert "LookinDoc DOCX Test" in doc.extracted_text
    assert "This paragraph contains natural language text" in doc.extracted_text

def test_extract_empty_file():
    with pytest.raises(DocumentExtractionError) as exc_info:
        extract_document(b"", "empty.txt")
    assert "empty" in str(exc_info.value).lower()

def test_extract_unsupported_format():
    with pytest.raises(DocumentExtractionError) as exc_info:
        extract_document(b"sample data", "image.png")
    assert "unsupported" in str(exc_info.value).lower()

def test_extract_pdf_corrupted():
    with pytest.raises(DocumentExtractionError) as exc_info:
        extract_document(b"this is not a valid pdf header", "corrupted.pdf")
    assert "failed to extract" in str(exc_info.value).lower() or "pdf" in str(exc_info.value).lower()
